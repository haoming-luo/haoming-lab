"""Create an editable Word edition from the same classroom content as PDF."""
import runpy,sys,tempfile,subprocess,re
from pathlib import Path
from lxml import etree
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, PageBreak, Image, Spacer
from reportlab.pdfgen.canvas import Canvas
from handout_schematics import ProblemDiagram

ROOT=Path(__file__).resolve().parents[1]
captured=[]
original=SimpleDocTemplate.build
SimpleDocTemplate.build=lambda self,story,**kwargs:captured.extend(story)
try:runpy.run_path(str(ROOT/'scripts/build-learning-pdf.py'))
finally:SimpleDocTemplate.build=original
doc=Document();sec=doc.sections[0]
sec.page_width=Pt(595.28);sec.page_height=Pt(841.89)
sec.left_margin=sec.right_margin=Pt(48)
sec.top_margin=Pt(48);sec.bottom_margin=Pt(45)
sec.header_distance=sec.footer_distance=Pt(22)
for name,font,size in [('Normal','Songti SC',10.5),('Title','Heiti SC',24),('Heading 1','Heiti SC',18),('Heading 2','Heiti SC',12),('Caption','Songti SC',8.6)]:
    s=doc.styles[name];s.font.name=font;s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
    s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=Pt(15)
    s.paragraph_format.space_before=Pt(0)
    s.font.bold=name.startswith('Heading')
    s.font.italic=False
    for border in s.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for fonts in s.element.xpath('.//w:rFonts'):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():del fonts.attrib[attr]
doc.core_properties.title='AI 辅助有限元分析'
doc.core_properties.subject='AgentFEM 基础实验讲义'
doc.core_properties.author='Haoming Luo'
hp=sec.header.paragraphs[0];hp.text='AgentFEM  /  基础实验讲义';hp.style='Caption'
fp=sec.footer.paragraphs[0];fp.style='Caption';fp.add_run('Haoming Luo');fp.paragraph_format.tab_stops.add_tab_stop(Pt(470))
fp.add_run('\t');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)

def runs(p,text):
    root=etree.fromstring(('<root>'+text.replace('<br/>','<br/>')+'</root>').encode())
    def walk(el,sup=None):
        if el.text:
            r=p.add_run(el.text)
            if sup=='super':r.font.superscript=True
            if sup=='sub':r.font.subscript=True
        for child in el:
            if child.tag=='br':p.add_run().add_break()
            elif child.tag=='link':
                h=OxmlElement('w:hyperlink');rid=p.part.relate_to(child.get('href'),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h.set(qn('r:id'),rid)
                r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'336B85');props.append(color);r.append(props);t=OxmlElement('w:t');t.text=''.join(child.itertext());r.append(t);h.append(r);p._p.append(h)
            else:walk(child,child.tag)
            if child.tail:p.add_run(child.tail)
    walk(root)

def para(obj,parent=doc,existing=None,table=False):
    st=obj.style;name=st.name
    style={'title':'Title','h1':'Heading 1','h2':'Heading 2','exercise_heading':'Heading 2','small':'Caption','reference':'Caption','label':'Heading 2'}.get(name,'Normal')
    p=existing if existing is not None else parent.add_paragraph()
    p.style=style;p.paragraph_format.space_after=Pt(3 if table else min(st.spaceAfter,8))
    p.paragraph_format.space_before=Pt(0 if table else min(st.spaceBefore,5))
    p.paragraph_format.line_spacing=Pt(12 if table else st.leading)
    p.paragraph_format.keep_with_next=name in ('title','h1','h2','exercise_heading','label')
    p.paragraph_format.widow_control=True
    runs(p,obj.text)
    for r in p.runs:
        r.font.size=Pt(9 if table else st.fontSize)
        r.font.name='Heiti SC' if st.fontName=='Hei' else 'Songti SC'
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),r.font.name)
    return p

tmp=Path(tempfile.mkdtemp(prefix='handout-word-'))
for item in captured:
    if isinstance(item,PageBreak):doc.add_page_break()
    elif isinstance(item,Paragraph):para(item)
    elif isinstance(item,Spacer):pass
    elif isinstance(item,Table):
        t=doc.add_table(rows=len(item._cellvalues),cols=len(item._colWidths));t.alignment=WD_TABLE_ALIGNMENT.LEFT;t.autofit=False
        for col,width in zip(t.columns,item._colWidths):col.width=Pt(width)
        props=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        props.append(borders)
        for row,values in zip(t.rows,item._cellvalues):
            for j,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=Pt(item._colWidths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                para(value,existing=cell.paragraphs[0],table=True)
                margins=OxmlElement('w:tcMar')
                for side,val in [('top','45'),('bottom','45'),('left','90'),('right','90')]:
                    m=OxmlElement('w:'+side);m.set(qn('w:w'),val);m.set(qn('w:type'),'dxa');margins.append(m)
                cell._tc.get_or_add_tcPr().append(margins)
        if item._cellvalues[0][0].style.name=='small':
            h=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(h)
            for cell in t.rows[0].cells:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(sh)
    elif isinstance(item,(Image,ProblemDiagram)):
        if isinstance(item,ProblemDiagram):
            pdf=tmp/f'diagram-{item.case}.pdf';c=Canvas(str(pdf),pagesize=(499,112));item.canv=c;item.draw();c.save()
            subprocess.run(['/Users/luo/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/bin/pdftoppm','-singlefile','-png','-r','220',str(pdf),str(pdf.with_suffix(''))],check=True)
            path=pdf.with_suffix('.png');width=499
        else:path=item.filename;width=item.drawWidth
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
        p.alignment=1;r=p.add_run();r.add_picture(str(path),width=Pt(width))
        pic=r._r.xpath('.//wp:docPr')[0];pic.set('descr',f'实验 {item.case} 模型与边界条件' if isinstance(item,ProblemDiagram) else Path(path).stem+' 参考结果图')

out=ROOT/'public/learn/AgentFEM-first-simulations.docx';doc.save(out);print(out)
